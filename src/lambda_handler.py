"""
Portfolio-safe Event Evaluation Lambda (placeholder version)

Pipeline:
S3 upload (PDF) -> Lambda
  -> download reference inputs from S3
  -> parse PDF + DOCX + TXT
  -> validate phrases/sections/pricing blocks
  -> generate Excel QA report
  -> upload report + copy PDF to output prefix
  -> (optional) notify Slack

This version:
- Uses generic prefixes (no internal project names)
- Avoids hardcoded proprietary phrases
- Uses tolerant matching to handle punctuation differences
- Avoids logging full event payloads
"""

import os
import io
import re
import json
import logging
from dataclasses import dataclass
from typing import Optional
from urllib.parse import unquote_plus
from datetime import datetime, timedelta, time as dtime
from textwrap import shorten

import boto3
import numpy as np
import pandas as pd

# Third-party libs (package via Lambda Layer or container image)
from docx import Document
from pypdf import PdfReader

# If you want Slack webhook (no token needed), use urllib (standard library)
import urllib.request


# =========================
# Logging
# =========================
logger = logging.getLogger()
logger.setLevel(os.environ.get("LOG_LEVEL", "INFO"))

s3 = boto3.client("s3")


# =========================
# ENV VARS (SAFE DEFAULTS)
# =========================
BUCKET            = os.environ.get("BUCKET")  # optional filter: process only this bucket if set
EVAL_PDFS_PREFIX  = os.environ.get("EVAL_PDFS_PREFIX", "eval_pdfs/")
REFERENCE_PREFIX  = os.environ.get("REFERENCE_PREFIX", "reference/")
PRICING_PREFIX    = os.environ.get("PRICING_PREFIX", "pricing/")
REPORTS_PREFIX    = os.environ.get("REPORTS_PREFIX", "reports/")
BASE_TMP          = os.environ.get("BASE_TMP", "/tmp/event_eval")

# Slack: webhook is safest for public repo examples
SLACK_WEBHOOK_URL = os.environ.get("SLACK_WEBHOOK_URL", "")

logger.info("Slack configured: %s", bool(SLACK_WEBHOOK_URL))


# =========================
# Slack helpers (webhook)
# =========================
def send_slack(text: str):
    if not SLACK_WEBHOOK_URL:
        logger.info("Slack not configured; skipping.")
        return

    payload = {"text": text}
    req = urllib.request.Request(
        SLACK_WEBHOOK_URL,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            logger.info("Slack webhook status: %s", resp.status)
    except Exception as e:
        logger.warning("Slack send failed: %s", e)


# =========================
# S3 helpers
# =========================
def s3_download(bucket: str, key: str, local_path: str):
    os.makedirs(os.path.dirname(local_path), exist_ok=True)
    logger.info("Downloading s3://%s/%s -> %s", bucket, key, local_path)
    s3.download_file(bucket, key, local_path)


def s3_upload(bucket: str, key: str, local_path: str, content_type: Optional[str] = None):
    logger.info("Uploading %s -> s3://%s/%s", local_path, bucket, key)
    extra = {}
    if content_type:
        extra["ContentType"] = content_type
    with open(local_path, "rb") as f:
        s3.put_object(Bucket=bucket, Key=key, Body=f.read(), **extra)


# =========================
# PDF helpers
# =========================
def _pdf_reader(pdf_path: str) -> PdfReader:
    return PdfReader(pdf_path)


def _pdf_page_texts(pdf_path: str) -> list[str]:
    reader = _pdf_reader(pdf_path)
    texts = []
    for p in reader.pages:
        t = p.extract_text() or ""
        texts.append(t)
    return texts


def _pdf_all_lines(pdf_path: str) -> list[str]:
    lines = []
    for t in _pdf_page_texts(pdf_path):
        lines.extend(t.splitlines())
    return lines


def _pdf_full_text(pdf_path: str) -> str:
    return "\n".join(_pdf_page_texts(pdf_path))


# =========================
# Normalization helpers
# (period/no period tolerant)
# =========================
def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", str(text)).strip()


def normalize_loose(text: str) -> str:
    """
    Loose normalization for comparisons:
    - collapse whitespace
    - lowercase
    - remove common punctuation so '.' ':' ',' don't cause false mismatches
    """
    s = normalize(text).lower()
    s = re.sub(r"[.,:;]", "", s)
    return s


# =========================
# Reference parsing
# =========================
def read_text_lines(path: str) -> list[str]:
    with open(path, encoding="utf-8") as f:
        return [line.strip() for line in f if line.strip()]


def extract_lines_from_docx(docx_path: str) -> list[str]:
    doc = Document(docx_path)
    return [normalize(p.text) for p in doc.paragraphs if p.text.strip()]


# =========================
# Phrase validation
# =========================
def verify_required_phrases_in_pdf(pdf_path: str, required_phrases: list[str]) -> pd.DataFrame:
    """
    Compare required phrases against PDF text using loose normalization.
    Never throws; returns Found/Not Found rows.
    """
    full_text_raw   = _pdf_full_text(pdf_path)
    full_text_loose = normalize_loose(full_text_raw)

    rows = []
    for phrase in required_phrases:
        phrase_loose = normalize_loose(phrase)
        found = phrase_loose in full_text_loose
        rows.append({
            "Expected Phrase": phrase[:200] + "..." if len(phrase) > 200 else phrase,
            "Result": "✅ Found" if found else "❌ Not Found"
        })

    return pd.DataFrame(rows)


# =========================
# Example pricing parsers (generic placeholders)
# =========================
def parse_pricing_txt(pricing_txt_path: str) -> pd.DataFrame:
    """
    Placeholder: parse a pricing txt into a structured DF.
    Keep your real parser here; this is generic.
    """
    with open(pricing_txt_path, encoding="utf-8") as f:
        lines = [ln.rstrip("\n") for ln in f if ln.strip()]

    # Example minimal structure
    return pd.DataFrame({"raw_line": lines})


def parse_pricing_blocks_from_pdf(pdf_path: str) -> pd.DataFrame:
    """
    Placeholder: parse pricing-like blocks from PDF.
    Keep your real parser here; this is generic.
    """
    lines = _pdf_all_lines(pdf_path)
    return pd.DataFrame({"raw_line": [ln.strip() for ln in lines if ln.strip()]})


def compare_pricing(pdf_df: pd.DataFrame, txt_df: pd.DataFrame) -> pd.DataFrame:
    """
    Placeholder comparison.
    Replace with your real side-by-side exact comparisons.
    """
    # Example output
    return pd.DataFrame([{
        "comparison": "placeholder",
        "mismatches": 0
    }])


# =========================
# Report builder (Excel)
# =========================
def build_eval_report(
    *,
    event_key: str,
    mapping_xlsx_path: str,
    reference_docx_path: str,
    reference_txt_path: str,
    pricing_txt_path: str,
    pdf_path: str,
    output_path: str,
):
    """
    Portfolio-safe report builder.
    Keeps same pattern: Summary + Pricing + Phrase Validation.
    """

    # --- Load mapping file (example)
    mapping_df = pd.read_excel(mapping_xlsx_path)

    # --- Example: get a "tier" (placeholder)
    tier_code = "TIER_X"
    if "Event_Name" in mapping_df.columns and "Tier_Code" in mapping_df.columns:
        match = mapping_df[mapping_df["Event_Name"].astype(str).str.strip() == str(event_key).strip()]
        if not match.empty:
            tier_code = str(match.iloc[0]["Tier_Code"]).strip()

    # --- Reference content
    docx_lines = extract_lines_from_docx(reference_docx_path)
    txt_lines  = read_text_lines(reference_txt_path)

    # --- Required phrases (SAFE)
    # In your real system, these come from a config file in S3 or repo.
    required_phrases = [
        "Sample required phrase A",
        "Sample required phrase B",
        "Sample required phrase C.",
        "Sample required phrase C",  # demonstrates punctuation tolerance
    ]
    phrase_df = verify_required_phrases_in_pdf(pdf_path, required_phrases)

    # --- Pricing (placeholders)
    pricing_txt_df = parse_pricing_txt(pricing_txt_path)
    pricing_pdf_df = parse_pricing_blocks_from_pdf(pdf_path)
    pricing_cmp_df = compare_pricing(pricing_pdf_df, pricing_txt_df)

    # --- Summary
    total_phrase_misses = int((phrase_df["Result"] == "❌ Not Found").sum())
    overall_status = "✅ All Good" if total_phrase_misses == 0 else "❌ Issues Found"

    summary_data = {
        "Event Key": event_key,
        "Tier": tier_code,
        "Total Phrase Misses": str(total_phrase_misses),
        "Overall Status": overall_status,
        "Notes": "This is a portfolio-safe sample. Replace placeholders with your real checks.",
    }
    summary_df = pd.DataFrame(summary_data.items(), columns=["Metric", "Value"])

    # --- Write Excel
    with pd.ExcelWriter(output_path, engine="xlsxwriter") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        phrase_df.to_excel(writer, sheet_name="Phrase Validation", index=False)
        pricing_cmp_df.to_excel(writer, sheet_name="Pricing Comparison", index=False)

        # Simple formatting
        wb = writer.book
        header_fmt = wb.add_format({"bold": True, "bg_color": "#404040", "font_color": "white"})
        left_align = wb.add_format({"align": "left"})

        for sheet_name, df_x in {
            "Summary": summary_df,
            "Phrase Validation": phrase_df,
            "Pricing Comparison": pricing_cmp_df,
        }.items():
            ws = writer.sheets[sheet_name]
            ws.freeze_panes(1, 0)
            for c, col in enumerate(df_x.columns):
                ws.write(0, c, col, header_fmt)

        # Summary column widths + left aligned values
        writer.sheets["Summary"].set_column("A:A", 35)
        writer.sheets["Summary"].set_column("B:B", 70, left_align)

    return overall_status, tier_code


# =========================
# Lambda handler
# =========================
def lambda_handler(event, context):
    # Avoid logging entire raw payload (can contain sensitive object keys)
    logger.info("Lambda invoked. Records=%s", len(event.get("Records", []) or []))

    try:
        records = event.get("Records", [])
        if not records:
            return {"statusCode": 400, "body": "No Records in event."}

        s3rec = records[0].get("s3", {})
        bucket = s3rec.get("bucket", {}).get("name")
        key = unquote_plus(s3rec.get("object", {}).get("key", ""))

        if not bucket or not key:
            return {"statusCode": 400, "body": "Missing bucket/key in event."}

        # Optional bucket filter
        if BUCKET and bucket != BUCKET:
            logger.info("Ignoring bucket %s (expected %s)", bucket, BUCKET)
            return {"statusCode": 200, "body": "Ignored (bucket filter)."}

        # Only process PDFs under inbound prefix
        if not key.startswith(EVAL_PDFS_PREFIX) or not key.lower().endswith(".pdf"):
            logger.info("Ignoring key: %s", key[:180])
            return {"statusCode": 200, "body": "Ignored (not inbound pdf)."}

        filename = os.path.basename(key)
        event_key = filename.rsplit(".", 1)[0]  # safe: "event key" derived from file name

        # Local paths
        work_dir = os.path.join(BASE_TMP, "work")
        ref_dir = os.path.join(BASE_TMP, "reference")
        pricing_dir = os.path.join(BASE_TMP, "pricing")
        os.makedirs(work_dir, exist_ok=True)
        os.makedirs(ref_dir, exist_ok=True)
        os.makedirs(pricing_dir, exist_ok=True)

        local_pdf_path = os.path.join(work_dir, filename)

        # ---- Download triggered PDF
        s3_download(bucket, key, local_pdf_path)

        # ---- Download reference inputs (generic names)
        mapping_xlsx = os.path.join(ref_dir, "EVENT_TIER_MAPPING.xlsx")
        reference_docx = os.path.join(ref_dir, "reference_section.docx")
        reference_txt  = os.path.join(ref_dir, "reference_phrases.txt")
        pricing_txt    = os.path.join(pricing_dir, "pricing_reference.txt")

        s3_download(bucket, f"{REFERENCE_PREFIX}EVENT_TIER_MAPPING.xlsx", mapping_xlsx)
        s3_download(bucket, f"{REFERENCE_PREFIX}reference_section.docx", reference_docx)
        s3_download(bucket, f"{REFERENCE_PREFIX}reference_phrases.txt", reference_txt)
        s3_download(bucket, f"{PRICING_PREFIX}pricing_reference.txt", pricing_txt)

        # ---- Build report
        output_filename = f"{event_key}_Eval_Comparison_Report.xlsx"
        local_output_path = os.path.join(work_dir, output_filename)

        if os.path.exists(local_output_path):
            os.remove(local_output_path)

        overall_status, tier_code = build_eval_report(
            event_key=event_key,
            mapping_xlsx_path=mapping_xlsx,
            reference_docx_path=reference_docx,
            reference_txt_path=reference_txt,
            pricing_txt_path=pricing_txt,
            pdf_path=local_pdf_path,
            output_path=local_output_path,
        )

        # ---- Upload outputs (avoid overwrite with timestamp)
        ts = datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
        report_dir_key = f"{REPORTS_PREFIX}tier_{str(tier_code).lower()}/{event_key}/"
        report_xlsx_key = f"{report_dir_key}{event_key}_Eval_Comparison_Report_{ts}.xlsx"
        report_pdf_key  = f"{report_dir_key}{event_key}_Eval_{ts}.pdf"

        s3_upload(bucket, report_xlsx_key, local_output_path, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        s3_upload(bucket, report_pdf_key, local_pdf_path, content_type="application/pdf")

        # (Optional) delete inbound object after processing
        delete_inbound = os.environ.get("DELETE_INBOUND", "true").lower() in ("1", "true", "yes")
        if delete_inbound:
            logger.info("Deleting inbound PDF: s3://%s/%s", bucket, key)
            s3.delete_object(Bucket=bucket, Key=key)

        # ---- Slack summary (no sensitive links)
        send_slack(
            f"📌 *Evaluation Completed*\n"
            f"File: `{filename}`\n"
            f"Tier: `{tier_code}`\n"
            f"Status: *{overall_status}*\n"
            f"Output prefix: `{report_dir_key}`"
        )

        return {"statusCode": 200, "body": f"Completed {event_key}: {overall_status}"}

    except Exception as e:
        logger.exception("Lambda failed")
        send_slack(f"❌ Eval Lambda error (portfolio-safe msg): {type(e).__name__}: {e}")
        return {"statusCode": 500, "body": f"{type(e).__name__}: {e}"}
